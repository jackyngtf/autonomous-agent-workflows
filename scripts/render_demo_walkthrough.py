"""Build a read-only, bilingual viewer from actual synthetic demo artifacts.

Run the success and blocked demos first, then:
    python scripts/render_demo_walkthrough.py
Serve the repository locally and open output/demo-walkthrough.html.
The viewer is a recording aid, not a production reporting application.
Requires only the standard library and the declared openpyxl dependency.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def load_workflow(root: Path, workflow: str) -> dict:
    """Check evidence against disk before displaying any success or safety claim."""
    records = {}
    for scenario in ("demo", "blocked"):
        folder = root / "output" / scenario / workflow
        evidence = json.loads((folder / "evidence.json").read_text(encoding="utf-8"))
        if evidence.get("synthetic") is not True or evidence.get("workflow") != workflow:
            raise ValueError(f"Expected synthetic {workflow} evidence in {scenario}")
        expected_status = evidence.get("status")
        allowed_statuses = {"PUBLISHED", "NO_CHANGE"} if scenario == "demo" else {"BLOCKED"}
        if expected_status not in allowed_statuses:
            raise ValueError(f"Unexpected {scenario} status: {expected_status}")
        source_before = evidence.get("input_sha256_before")
        source_after = evidence.get("input_sha256_after")
        if not source_before or source_before != source_after or evidence.get("input_unchanged") is not True:
            raise ValueError(f"Source preservation is not established for {scenario}/{workflow}")
        entries = evidence.get("workbooks", [])
        if len(entries) != 1:
            raise ValueError("This short walkthrough expects exactly one workbook per workflow")
        entry = entries[0]
        if entry.get("workbook") != "2026-08.xlsx" or entry.get("status") != expected_status:
            raise ValueError("Unexpected workbook or status")
        if entry.get("cleanup_executed") is not False:
            raise ValueError("This recording must not imply that cleanup ran")
        workbook_path = folder / entry["workbook"]
        actual_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
        if actual_hash != entry.get("workbook_sha256_after"):
            raise ValueError(f"Workbook has changed since evidence capture: {workbook_path}")
        if scenario == "blocked":
            if (entry.get("workbook_sha256_before") != actual_hash
                    or entry.get("cleanup_eligible_dates") != []
                    or "content mismatch" not in entry.get("reason", "").lower()):
                raise ValueError("Blocked evidence does not establish the depicted preservation outcome")
        records[scenario] = evidence

    workbook = openpyxl.load_workbook(
        root / "output/demo" / workflow / "2026-08.xlsx", read_only=True, data_only=False)
    try:
        sheets = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                raise ValueError("Unexpected empty worksheet")
            columns = [0, 3, 4, 5, 12] if workflow == "avaya" else [0, 2, 3, 5, 7]
            headers = [str(rows[0][index] or "") for index in columns]
            preview = [[str(row[index] or "") for index in columns] for row in rows[1:7]]
            sheets.append({"name": sheet.title, "row_count": len(rows) - 1,
                           "column_count": len(rows[0]), "headers": headers, "rows": preview})
    finally:
        workbook.close()
    reported = records["demo"]["workbooks"][0]
    if sum(sheet["row_count"] for sheet in sheets) != reported.get("total_rows"):
        raise ValueError("Workbook row total disagrees with evidence")
    return {"workflow": workflow, "sheets": sheets, "evidence": records}


PAGE = r'''<!doctype html>
<html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Synthetic demo · artifact walkthrough</title><link rel="icon" href="data:,">
<style>
*{box-sizing:border-box}body{margin:0;background:#f3f5f4;color:#152b29;font:16px/1.45 Arial,"Microsoft JhengHei",sans-serif}
main{max-width:1370px;margin:auto;padding:28px 38px}header{display:flex;justify-content:space-between;align-items:start;gap:24px}
.eyebrow{font-size:12px;letter-spacing:1.7px;color:#526562;font-weight:700;text-transform:uppercase}h1{font-size:32px;letter-spacing:-.8px;margin:7px 0 3px}h2{font-size:23px;margin:0}p{margin:5px 0;color:#526562}.scope{background:#e3eeea;border:1px solid #c5d9d1;padding:12px 17px;border-radius:7px;font-size:14px;min-width:285px}.scope strong{display:block;color:#12634e}
nav{display:flex;gap:8px;margin:24px 0 18px}button{font:inherit;cursor:pointer;border:1px solid #cad5d0;background:white;color:#314a44;border-radius:6px;padding:10px 18px}button[aria-pressed=true]{background:#164f42;color:white;border-color:#164f42}button:focus-visible{outline:3px solid #d99d3d;outline-offset:3px}.panel{background:white;border:1px solid #d5dfda;border-radius:9px;padding:25px 27px;min-height:580px}.panel-title{display:flex;justify-content:space-between;align-items:start}.badge{font-size:12px;font-weight:bold;letter-spacing:.7px;padding:6px 10px;background:#e4f0e9;color:#1f644a;border-radius:4px}.muted{color:#62736e}.small{font-size:13px}.metrics{display:flex;border-block:1px solid #e4e9e6;margin:20px 0 17px;padding:14px 0}.metric{min-width:175px;padding:0 24px;border-right:1px solid #e4e9e6}.metric:first-child{padding-left:0}.metric:last-child{border:0}.metric strong{display:block;font-size:25px;font-weight:600}.metric span{font-size:13px;color:#62736e}.tabs{display:flex;gap:7px;margin:16px 0 13px}.tabs button{font-size:14px;padding:7px 12px}table{width:100%;border-collapse:collapse;font-size:13px;table-layout:fixed}th{text-align:left;background:#f0f4f1;padding:11px 12px;font-weight:600}td{padding:11px 12px;border-bottom:1px solid #edf0ee;overflow-wrap:anywhere;vertical-align:top}td:first-child,th:first-child{width:24%}.table-note{margin-top:10px;font-size:12px;color:#62736e}.proof{background:#f3f7f4;border-left:3px solid #43876c;padding:12px 15px;margin-top:19px;font-size:14px}.proof strong{display:block}.footer{display:flex;justify-content:space-between;margin-top:16px;font-size:12px;color:#66766f}.blocked{background:#f7e9df;color:#994d27}.block-grid{display:grid;grid-template-columns:1fr 1fr;gap:20px;margin-top:23px}.block-card{border:1px solid #dbe2de;border-radius:7px;padding:19px}.block-card h3{margin:0 0 14px;font-size:18px}.fact{display:flex;justify-content:space-between;gap:20px;border-bottom:1px solid #e8edea;padding:10px 0;font-size:14px}.fact strong{color:#196546}.hash{font:12px/1.7 Consolas,monospace;word-break:break-all;background:#f4f6f4;padding:10px;margin-top:8px;color:#40544c}.reason{padding:12px 14px;background:#fbf2e9;border-left:3px solid #bf8157;margin-top:20px;font-size:14px}.path{font:12px Consolas,monospace;color:#617069}.step{font-size:12px;color:#65756f;margin-bottom:6px}
</style>
<main>
<header><div><div class="eyebrow">Autonomous Agent Workflows · Portfolio evidence</div><h1>From records to a checked workbook</h1><p>從來源紀錄到已驗證工作簿 · Read-only artifact walkthrough</p></div><div class="scope"><strong>SYNTHETIC · LOCAL SIMULATION</strong>合成資料 · 本機模擬<br>No NAS connection or source deletion<br>不連接 NAS，不刪除來源</div></header>
<nav aria-label="Artifact view"><button id="avaya" aria-pressed="true">01 · Avaya workbook／通話月報</button><button id="nas" aria-pressed="false">02 · NAS workbook／存取月報</button><button id="blocked" aria-pressed="false">03 · Blocked update／阻止更新</button></nav>
<section class="panel" id="panel" aria-live="polite"></section>
<div class="footer"><span>Values read from generated XLSX + evidence.json · 數值來自實際產出</span><span>Viewer only · Not a production application／並非正式應用程式</span></div>
</main>
<script>
const data = __DATA__;
const panel = document.getElementById('panel');
function element(tag, text, cls){const el=document.createElement(tag);if(text!==undefined)el.textContent=text;if(cls)el.className=cls;return el;}
function metric(parent,value,label){const el=element('div',undefined,'metric');el.append(element('strong',String(value)),element('span',label));parent.append(el);}
function choose(view){for(const button of document.querySelectorAll('nav button'))button.setAttribute('aria-pressed',String(button.id===view));panel.replaceChildren();if(view==='blocked'){showBlocked();return;}showWorkbook(view);}
function showWorkbook(name){const item=data[name], ev=item.evidence.demo, result=ev.workbooks[0];const top=element('div',undefined,'panel-title');const title=element('div');title.append(element('div','GENERATED WORKBOOK／實際產生的工作簿','step'),element('h2',name==='avaya'?'Avaya · Monthly call log':'NAS · Monthly file-transfer log'),element('p','output/demo/'+name+'/2026-08.xlsx','path'));top.append(title,element('span',result.status==='PUBLISHED'?'PUBLISHED LOCALLY／已本機發佈':'NO CHANGE／沒有變更','badge'));panel.append(top);
const stats=element('div',undefined,'metrics');metric(stats,result.new_rows,'New records／新增紀錄');metric(stats,result.preserved_rows,'Preserved records／保留紀錄');metric(stats,item.sheets.length,'Worksheets／工作表');metric(stats,'0','Sources deleted／來源刪除');panel.append(stats);
const tabs=element('div',undefined,'tabs'), tableArea=element('div');panel.append(tabs,tableArea);
function sheetView(sheet){for(const b of tabs.children)b.setAttribute('aria-pressed',String(b.dataset.sheet===sheet.name));tableArea.replaceChildren();const table=element('table'),head=element('thead'),hr=element('tr');for(const value of sheet.headers)hr.append(element('th',value));head.append(hr);table.append(head);const body=element('tbody');for(const values of sheet.rows){const row=element('tr');for(const value of values)row.append(element('td',value));body.append(row);}table.append(body);tableArea.append(table,element('div',`Showing ${sheet.rows.length} of ${sheet.row_count} data rows · selected ${sheet.headers.length} of ${sheet.column_count} columns／顯示部分欄位；全部儲存格另經程式驗證`,'table-note'));}
for(const sheet of item.sheets){const added=result.new_sheets.includes(sheet.name),button=element('button',sheet.name+(added?' · added／新增':' · preserved／保留'));button.dataset.sheet=sheet.name;button.onclick=()=>sheetView(sheet);tabs.append(button);}sheetView(item.sheets[item.sheets.length-1]);
const proof=element('div',undefined,'proof');proof.append(element('strong','Verification recorded／已記錄的驗證'),element('span',result.status==='PUBLISHED'?'Every cell checked locally; simulated received copy checked by cell values and SHA-256.／本機逐格檢查，模擬取回副本再比對內容及雜湊。':'Existing workbook read; supplied dates checked for conflicts. No new publication on this run.／讀取既有工作簿並檢查日期衝突；這次執行沒有重新發佈。'));panel.append(proof,element('p','Fixture clock／固定示範時間：'+ev.fixture_clock,'small muted'));}
function showBlocked(){const top=element('div',undefined,'panel-title'),title=element('div');title.append(element('div','INJECTED CONTENT MISMATCH／刻意加入內容差異','step'),element('h2','The received copy fails verification'),element('p','模擬取回副本驗證失敗；更新被阻止'));top.append(title,element('span','BLOCKED／已阻止','badge blocked'));panel.append(top);const cards=element('div',undefined,'block-grid');for(const name of ['avaya','nas']){const ev=data[name].evidence.blocked, result=ev.workbooks[0],card=element('div',undefined,'block-card');card.append(element('h3',name==='avaya'?'Avaya call log／通話紀錄':'NAS access log／存取紀錄'));for(const [label,value] of [['Destination unchanged／目的地未變','SHA-256 MATCH'],['Source unchanged／來源未變','TRUE'],['Cleanup performed／已執行清理','FALSE']]){const row=element('div',undefined,'fact');row.append(element('span',label),element('strong',value));card.append(row);}card.append(element('p','Destination SHA-256 before = after／目的地前後雜湊相同','small'),element('div',result.workbook_sha256_before,'hash'),element('p','output/blocked/'+name+'/evidence.json','path'));cards.append(card);}panel.append(cards);const reason=element('div',undefined,'reason');reason.append(element('strong','Recorded reason／記錄原因：'),element('div',data.avaya.evidence.blocked.workbooks[0].reason));panel.append(reason);panel.append(element('div','This local demo checks a simulated received copy before replacing the destination. It does not prove atomic SMB upload or production rollback.／本機示範先檢查模擬副本，再替換目的地；不代表正式 SMB 上傳或回復保證。','proof'));}
for(const name of ['avaya','nas','blocked'])document.getElementById(name).onclick=()=>choose(name);choose('avaya');
</script></html>'''


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "output/demo-walkthrough.html")
    args = parser.parse_args()
    data = {name: load_workflow(ROOT, name) for name in ("avaya", "nas")}
    payload = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(PAGE.replace("__DATA__", payload), encoding="utf-8")
    print(args.output)


if __name__ == "__main__":
    main()
