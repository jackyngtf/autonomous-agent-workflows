"""Deterministic reporting cores with input and publication gates.

Only fixed-schema text reporting workbooks are supported. This is not a
general Excel editor, network adapter, or transactional storage system.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Callable
from zoneinfo import ZoneInfo

import openpyxl
import xlsxwriter

MELBOURNE = ZoneInfo("Australia/Melbourne")
SMDR_HEADERS = (
    "Call Start", "Connected Time", "Ring Time", "Caller", "Direction",
    "Called Number", "Dialled Number", "Account", "Is Internal", "Call ID",
    "Continuation", "Party1Device", "Party1Name", "Party2Device", "Party2Name",
    "Hold Time", "Park Time", "AuthValid", "AuthCode", "UserCharged",
    "CallCharge", "Currency", "AmountAtLastUserChange", "CallUnits",
    "UnitsAtLastUserChange", "CostPerUnit", "MarkUp", "ExternalTargetingCause",
    "ExternalTargeterId", "ExternalTargetedNumber",
)
NAS_SOURCE_HEADERS = ("UTC", "Account", "Event", "Share", "Path", "File size", "Action", "From")
NAS_HEADERS = ("Local Time (Australia/Melbourne)",) + NAS_SOURCE_HEADERS
Rows = tuple[tuple[str, ...], ...]
Sheets = dict[str, Rows]


class ValidationError(ValueError):
    """An input or workbook failed a check; publication must stop."""


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def local_day(instant: datetime) -> date:
    if instant.tzinfo is None or instant.utcoffset() is None:
        raise ValidationError("An explicit timezone is required.")
    return instant.astimezone(MELBOURNE).date()


def read_avaya(path: Path) -> dict[date, Rows]:
    """Require the complete 30-column schema and validate every data row."""
    match = re.fullmatch(r"smdr_(\d{4}-\d{2}-\d{2})\.csv", path.name)
    if not match:
        raise ValidationError("SMDR filename must contain an ISO date.")
    try:
        file_day = date.fromisoformat(match[1])
        text = path.read_text(encoding="utf-8-sig")
        if "\x00" in text:
            raise ValidationError("SMDR input contains NUL bytes.")
        parsed = list(csv.reader(io.StringIO(text), strict=True))
        if not parsed or tuple(parsed[0]) != SMDR_HEADERS:
            raise ValidationError("SMDR header does not match the 30-column schema.")
        if len(parsed) < 2:
            raise ValidationError("SMDR input has no data rows.")
        rows = []
        for line, row in enumerate(parsed[1:], 2):
            if len(row) != len(SMDR_HEADERS):
                raise ValidationError(f"SMDR row {line}: expected 30 columns, found {len(row)}.")
            stamp = datetime.strptime(row[0].replace("/", "-"), "%Y-%m-%d %H:%M:%S")
            if stamp.date() != file_day:
                raise ValidationError(f"SMDR row {line}: date differs from filename.")
            row[0] = stamp.isoformat(sep=" ")
            rows.append(tuple(row))
        return {file_day: tuple(rows)}
    except (UnicodeError, csv.Error, ValueError) as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError(f"Invalid SMDR input: {exc}") from exc


def read_nas(path: Path, result_limit: int = 50_000) -> dict[date, Rows]:
    """A result reaching the query cap has unproven completeness and is blocked."""
    try:
        records = json.loads(path.read_text(encoding="utf-8-sig"))
        if not isinstance(records, list) or not records:
            raise ValidationError("NAS input must be a nonempty list of records.")
        if len(records) >= result_limit:
            raise ValidationError("NAS result reached the query cap; completeness is unproven.")
        grouped: dict[date, list[tuple[str, ...]]] = {}
        for number, record in enumerate(records, 1):
            if not isinstance(record, dict) or set(record) != set(NAS_SOURCE_HEADERS):
                raise ValidationError(f"NAS record {number}: unexpected schema.")
            if not all(isinstance(value, str) for value in record.values()):
                raise ValidationError(f"NAS record {number}: all fields must be strings.")
            stamp = datetime.fromisoformat(record["UTC"].replace("Z", "+00:00"))
            day = local_day(stamp)
            if int(record["File size"]) < 0:
                raise ValidationError(f"NAS record {number}: negative file size.")
            row = (stamp.astimezone(MELBOURNE).isoformat(),) + tuple(record[key] for key in NAS_SOURCE_HEADERS)
            grouped.setdefault(day, []).append(row)
        return {day: tuple(rows) for day, rows in sorted(grouped.items())}
    except (UnicodeError, ValueError, TypeError) as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError(f"Invalid NAS input: {exc}") from exc


def completed_dates(groups: dict[date, Rows], now: datetime) -> dict[date, Rows]:
    """Include completed weekends; exclude today and the future in Melbourne."""
    today = local_day(now)
    return {day: rows for day, rows in groups.items() if day < today}


def by_month(groups: dict[date, Rows]) -> dict[str, Sheets]:
    months: dict[str, Sheets] = {}
    for day, rows in sorted(groups.items()):
        months.setdefault(day.strftime("%Y-%m"), {})[day.isoformat()] = rows
    return months


def read_workbook(path: Path, headers: tuple[str, ...]) -> Sheets:
    """Only FileNotFoundError means absent; an existing unreadable file stops work."""
    try:
        path.stat()
    except FileNotFoundError:
        return {}
    try:
        # These demo workbooks are small. Full loading reads the actual cell
        # records rather than trusting potentially stale worksheet dimensions.
        workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
        try:
            sheets = {}
            for sheet in workbook.worksheets:
                if sheet.max_column != len(headers):
                    raise ValidationError(f"{sheet.title}: unexpected column count.")
                values = []
                for cells in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValidationError("Formula-bearing workbooks are outside this demo's scope.")
                    if any(cell.value is not None and not isinstance(cell.value, str) for cell in cells):
                        raise ValidationError("Only text reporting cells are supported.")
                    values.append(tuple("" if cell.value is None else cell.value for cell in cells))
                if not values or values[0] != headers:
                    raise ValidationError(f"{sheet.title}: header mismatch.")
                sheets[sheet.title] = tuple(values[1:])
            return sheets
        finally:
            workbook.close()
    except Exception as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError(f"Existing workbook is unreadable: {path.name} ({type(exc).__name__}).") from exc


def write_workbook(path: Path, headers: tuple[str, ...], sheets: Sheets) -> None:
    if not sheets:
        raise ValidationError("Refusing to create a workbook without data sheets.")
    with xlsxwriter.Workbook(path, {"strings_to_formulas": False, "strings_to_urls": False}) as workbook:
        workbook.set_properties({"title": "Synthetic reporting demo", "created": datetime(2026, 8, 5)})
        heading = workbook.add_format({"bold": True, "bg_color": "#153D3A", "font_color": "#FFFFFF", "text_wrap": True})
        normal = workbook.add_format({"valign": "top"})
        for name, rows in sorted(sheets.items()):
            sheet = workbook.add_worksheet(name)
            sheet.write_row(0, 0, headers, heading)
            sheet.set_row(0, 32)
            sheet.set_column(0, len(headers) - 1, 18)
            sheet.set_column(0, 0, 28)
            for index, row in enumerate(rows, 1):
                if len(row) != len(headers):
                    raise ValidationError(f"{name}: row width does not match headers.")
                sheet.write_row(index, 0, row, normal)
            sheet.freeze_panes(1, 0)
            sheet.autofilter(0, 0, len(rows), len(headers) - 1)


def verify_workbook(path: Path, headers: tuple[str, ...], expected: Sheets) -> None:
    """Read every header/data cell using a different library than the writer."""
    if read_workbook(path, headers) != expected:
        raise ValidationError("Workbook content mismatch: sheet names, row counts or cell values differ.")


def update_workbook(destination: Path, headers: tuple[str, ...], incoming: Sheets,
                    after_transfer: Callable[[Path], None] | None = None) -> dict:
    """Publish the verified local transfer candidate; never remove source files."""
    existing = read_workbook(destination, headers)
    for name in existing.keys() & incoming.keys():
        if existing[name] != incoming[name]:
            raise ValidationError(f"Existing date {name} conflicts with the supplied input.")
    added = {name: rows for name, rows in incoming.items() if name not in existing}
    expected = {**existing, **added}
    before = digest(destination) if destination.exists() else None
    result = {
        "status": "NO_CHANGE", "new_sheets": sorted(added), "new_rows": sum(map(len, added.values())),
        "preserved_sheets": len(existing), "preserved_rows": sum(map(len, existing.values())),
        "total_rows": sum(map(len, expected.values())), "cleanup_eligible_dates": [], "cleanup_executed": False,
        "workbook_sha256_before": before, "workbook_sha256_after": before,
        "checks": ["existing workbook readable", "existing date conflicts checked"],
    }
    if not added:
        return result
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".demo-stage-", dir=destination.parent) as scratch:
        candidate, received = Path(scratch) / "candidate.xlsx", Path(scratch) / "received.xlsx"
        write_workbook(candidate, headers, expected)
        verify_workbook(candidate, headers, expected)
        result["checks"].append("local candidate: every cell verified with openpyxl")
        shutil.copyfile(candidate, received)
        if after_transfer:
            after_transfer(received)
        verify_workbook(received, headers, expected)
        if digest(candidate) != digest(received):
            raise ValidationError("Simulated transfer digest mismatch.")
        result["checks"].append("simulated received copy: every cell and SHA-256 verified")
        # An intervening-writer check, not a cross-process lock or network transaction.
        current = digest(destination) if destination.exists() else None
        if current != before:
            raise ValidationError("Destination changed during the run.")
        os.replace(received, destination)
    result.update(status="PUBLISHED", cleanup_eligible_dates=sorted(added), workbook_sha256_after=digest(destination))
    return result


def tamper_received_copy(path: Path) -> None:
    """Negative showcase: same dimensions, changed cell after simulated transfer."""
    workbook = openpyxl.load_workbook(path)
    try:
        workbook.worksheets[0]["A2"] = "SIMULATED TAMPER / 模擬內容改動"
        workbook.save(path)
    finally:
        workbook.close()


def seed_sheet(workflow: str) -> Sheets:
    """Fictional prior-day record independent of incoming fixtures."""
    if workflow == "avaya":
        row = [""] * len(SMDR_HEADERS)
        row[:10] = ["2026-08-03 09:00:00", "00:01:00", "00:00:04", "EXT-101", "I", "EXT-201", "EXT-201", "DEMO", "1", "DEMO-PREV-001"]
    else:
        row = ["2026-08-03T09:00:00+10:00", "2026-08-02T23:00:00+00:00", "operator-alpha", "file-transfer", "demo", "/demo/prior-day.txt", "128", "read", "192.0.2.10 (DEMO-01)"]
    return {"2026-08-03": (tuple(row),)}


def write_report(output: Path, evidence: dict) -> None:
    output.mkdir(parents=True, exist_ok=True)
    (output / "evidence.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = [f"# {evidence['workflow'].upper()} — synthetic run / 合成資料執行報告", "",
        "Offline demonstration only. No production connection or source deletion.",
        "只供離線示範；沒有連接正式系統，亦沒有刪除來源檔案。", "",
        f"Fixture clock / 固定示範時間: `{evidence['fixture_clock']}`", "",
        "| Check / 檢查 | Result / 結果 |", "|---|---|",
        f"| Outcome / 執行結果 | {evidence['status']} |",
        f"| Input unchanged / 來源維持不變 | {evidence['input_unchanged']} |",
        "| Cleanup executed / 已執行清理 | False |", ""]
    for item in evidence["workbooks"]:
        lines += [f"## {item['workbook']}", "", f"Status / 狀態: **{item['status']}**", ""]
        if item["status"] == "BLOCKED":
            lines += [f"Reason / 原因: {item['reason']}", "",
                      "Publication blocked; prior workbook retained. / 已阻止發佈，保留原有工作簿。", ""]
        else:
            lines += [f"Added / 新增: {item['new_rows']} rows / 筆, {len(item['new_sheets'])} sheets / 個工作表.",
                      f"Preserved / 保留: {item['preserved_rows']} rows / 筆, {item['preserved_sheets']} sheets / 個工作表.",
                      f"Total / 總數: {item['total_rows']} rows / 筆.", ""]
        lines += ["Cleanup eligibility is recorded only; no input is removed. / 只記錄清理資格，不會移除輸入。", ""]
    lines += ["Verification reads every text cell with openpyxl against the pre-write snapshot and validated input.",
              "驗證以 openpyxl 讀回所有文字儲存格，與寫入前快照及已驗證輸入逐一比較。", "",
              "See evidence.json for checks, hashes and scope. / 檢查項目、雜湊及範圍詳見 evidence.json。", ""]
    (output / "run-report.md").write_text("\n".join(lines), encoding="utf-8")


def run_demo(workflow: str, source: Path, output: Path, now: datetime, blocked: bool = False) -> dict:
    headers = SMDR_HEADERS if workflow == "avaya" else NAS_HEADERS
    input_before = digest(source)
    evidence = {"workflow": workflow, "fixture_clock": now.isoformat(), "synthetic": True,
                "scope": "Local transfer simulation; text cell preservation; no network or input deletion.",
                "status": "NO_CHANGE", "workbooks": [], "input_sha256_before": input_before}
    output.mkdir(parents=True, exist_ok=True)
    try:
        groups = read_avaya(source) if workflow == "avaya" else read_nas(source)
        planned = completed_dates(groups, now)
        evidence["deferred_dates"] = [day.isoformat() for day in sorted(groups.keys() - planned.keys())]
        for month, incoming in by_month(planned).items():
            destination = output / f"{month}.xlsx"
            # Bootstrap the fixture's prior-day workbook only on confirmed absence.
            try:
                destination.stat()
            except FileNotFoundError:
                if month == "2026-08":
                    write_workbook(destination, headers, seed_sheet(workflow))
                    verify_workbook(destination, headers, seed_sheet(workflow))
            before = digest(destination) if destination.exists() else None
            try:
                result = update_workbook(destination, headers, incoming, tamper_received_copy if blocked else None)
                result["workbook"] = destination.name
                evidence["workbooks"].append(result)
            except ValidationError as exc:
                evidence["workbooks"].append({"workbook": destination.name, "status": "BLOCKED", "reason": str(exc),
                    "workbook_sha256_before": before, "workbook_sha256_after": digest(destination) if destination.exists() else None,
                    "cleanup_eligible_dates": [], "cleanup_executed": False})
    except (ValidationError, OSError) as exc:
        evidence["workbooks"].append({"workbook": "input-validation", "status": "BLOCKED", "reason": str(exc),
                                      "cleanup_eligible_dates": [], "cleanup_executed": False})
    statuses = {item["status"] for item in evidence["workbooks"]}
    evidence["status"] = "BLOCKED" if "BLOCKED" in statuses else "PUBLISHED" if "PUBLISHED" in statuses else "NO_CHANGE"
    evidence["input_sha256_after"] = digest(source)
    evidence["input_unchanged"] = evidence["input_sha256_after"] == input_before
    write_report(output, evidence)
    return evidence
