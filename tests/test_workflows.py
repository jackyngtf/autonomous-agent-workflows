"""Regression checks for realistic reporting failure modes; no network access."""
import csv
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import patch

import openpyxl

from demo.__main__ import FIXTURES, ROOT
from demo.workflows import (
    MELBOURNE, NAS_HEADERS, SMDR_HEADERS, ValidationError, by_month,
    completed_dates, digest, local_day, read_avaya, read_nas, read_workbook,
    run_demo, seed_sheet, tamper_received_copy, update_workbook,
    verify_workbook, write_workbook,
)

NOW = datetime(2026, 8, 5, 10, tzinfo=MELBOURNE)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory()
        self.addCleanup(self.scratch.cleanup)
        self.folder = Path(self.scratch.name)

    def csv_copy(self, change):
        with FIXTURES["avaya"].open(newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
        change(rows)
        path = self.folder / FIXTURES["avaya"].name
        with path.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerows(rows)
        return path

    def test_fixture_schema_and_known_literal_values(self):
        self.assertEqual(len(SMDR_HEADERS), 30)
        avaya = read_avaya(FIXTURES["avaya"])[date(2026, 8, 4)]
        nas = read_nas(FIXTURES["nas"])[date(2026, 8, 4)]
        self.assertEqual(len(avaya), 13)
        self.assertTrue(all(len(row) == 30 for row in avaya))
        self.assertEqual(avaya[0][:4], ("2026-08-04 09:00:00", "00:01:30", "00:00:04", "EXT-101"))
        self.assertEqual(len(nas), 8)
        self.assertEqual(nas[0][0], "2026-08-04T10:05:00+10:00")
        self.assertEqual(nas[0][2], "anonymous")  # Never infer identity from client strings.

    def test_both_workflows_preserve_existing_cells_and_rerun_without_writing(self):
        for name, headers, added in (("avaya", SMDR_HEADERS, 13), ("nas", NAS_HEADERS, 8)):
            with self.subTest(workflow=name):
                output = self.folder / name
                result = run_demo(name, FIXTURES[name], output, NOW)
                self.assertEqual(result["status"], "PUBLISHED")
                self.assertTrue(result["input_unchanged"])
                self.assertEqual(result["workbooks"][0]["new_rows"], added)
                workbook = output / "2026-08.xlsx"
                contents = read_workbook(workbook, headers)
                self.assertEqual(contents["2026-08-03"], seed_sheet(name)["2026-08-03"])
                self.assertEqual(len(contents["2026-08-04"]), added)
                before = digest(workbook)
                repeat = run_demo(name, FIXTURES[name], output, NOW)
                self.assertEqual(repeat["status"], "NO_CHANGE")
                self.assertEqual(digest(workbook), before)
                self.assertEqual(repeat["workbooks"][0]["new_rows"], 0)

    def test_schema_mismatch_and_mixed_probe_rows_are_rejected(self):
        changes = (
            lambda rows: rows[0].__setitem__(0, "Wrong header"),
            lambda rows: rows[1].append("extra"),
            lambda rows: rows.append(["GET / HTTP/1.0"]),
            lambda rows: rows[1].__setitem__(3, "bad\x00value"),
            lambda rows: rows[1].__setitem__(0, "2026/08/05 09:00:00"),
            lambda rows: rows[1].__setitem__(0, "2026/08/04 25:99:00"),
        )
        for change in changes:
            with self.subTest(change=change):
                with self.assertRaises(ValidationError):
                    read_avaya(self.csv_copy(change))

    def test_invalid_input_writes_blocked_report_without_a_workbook(self):
        source = self.csv_copy(lambda rows: rows.append(["GET / HTTP/1.0"]))
        result = run_demo("avaya", source, self.folder / "out", NOW)
        self.assertEqual(result["status"], "BLOCKED")
        self.assertFalse((self.folder / "out/2026-08.xlsx").exists())
        self.assertTrue(result["input_unchanged"])
        self.assertIn("合成資料", (self.folder / "out/run-report.md").read_text(encoding="utf-8"))

    def test_absence_and_unreadable_workbook_are_not_conflated(self):
        path = self.folder / "bad.xlsx"
        self.assertEqual(read_workbook(path, SMDR_HEADERS), {})
        path.write_bytes(b"an existing damaged workbook")
        before = digest(path)
        with self.assertRaisesRegex(ValidationError, "unreadable"):
            update_workbook(path, SMDR_HEADERS, seed_sheet("avaya"))
        self.assertEqual(digest(path), before)
        with patch.object(Path, "stat", side_effect=PermissionError("denied")):
            with self.assertRaises(PermissionError):
                read_workbook(path, SMDR_HEADERS)

    def test_conflicting_existing_date_blocks_replacement(self):
        path = self.folder / "existing.xlsx"
        existing = seed_sheet("avaya")
        write_workbook(path, SMDR_HEADERS, existing)
        before = digest(path)
        altered = list(existing["2026-08-03"][0])
        altered[3] = "DIFFERENT CALLER"
        with self.assertRaisesRegex(ValidationError, "conflicts"):
            update_workbook(path, SMDR_HEADERS, {"2026-08-03": (tuple(altered),)})
        self.assertEqual(digest(path), before)

    def test_understated_worksheet_dimension_cannot_hide_existing_rows(self):
        path = self.folder / "understated.xlsx"
        first = seed_sheet("avaya")["2026-08-03"][0]
        second = list(first)
        second[9] = "DEMO-PREV-002"
        existing = {"2026-08-03": (first, tuple(second))}
        write_workbook(path, SMDR_HEADERS, existing)
        # Alter only the advertised bounds; the second data row remains in XML.
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        parts["xl/worksheets/sheet1.xml"] = parts["xl/worksheets/sheet1.xml"].replace(
            b'<dimension ref="A1:AD3"/>', b'<dimension ref="A1:AD2"/>')
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for name, content in parts.items():
                archive.writestr(name, content)
        incoming = {"2026-08-04": read_avaya(FIXTURES["avaya"])[date(2026, 8, 4)]}
        result = update_workbook(path, SMDR_HEADERS, incoming)
        self.assertEqual(result["preserved_rows"], 2)
        actual = read_workbook(path, SMDR_HEADERS)
        self.assertEqual(actual["2026-08-03"], existing["2026-08-03"])
        self.assertEqual(actual["2026-08-04"], incoming["2026-08-04"])

    def test_same_dimensions_tamper_blocks_publication_and_cleanup(self):
        for name in FIXTURES:
            with self.subTest(workflow=name):
                result = run_demo(name, FIXTURES[name], self.folder / name, NOW, blocked=True)
                self.assertEqual(result["status"], "BLOCKED")
                report = result["workbooks"][0]
                self.assertIn("content mismatch", report["reason"])
                self.assertEqual(report["workbook_sha256_before"], report["workbook_sha256_after"])
                self.assertEqual(report["cleanup_eligible_dates"], [])
                self.assertFalse(report["cleanup_executed"])
                self.assertTrue(result["input_unchanged"])

    def test_duplicate_header_regression_fails_independent_verification(self):
        path = self.folder / "duplicate.xlsx"
        expected = seed_sheet("nas")
        write_workbook(path, NAS_HEADERS, expected)
        workbook = openpyxl.load_workbook(path)
        workbook.active.insert_rows(2)
        for column, value in enumerate(NAS_HEADERS, 1):
            workbook.active.cell(2, column, value)
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValidationError, "content mismatch"):
            verify_workbook(path, NAS_HEADERS, expected)

    def test_partial_transfer_cannot_replace_destination(self):
        path = self.folder / "partial.xlsx"
        write_workbook(path, SMDR_HEADERS, seed_sheet("avaya"))
        before = digest(path)
        incoming = {"2026-08-04": read_avaya(FIXTURES["avaya"])[date(2026, 8, 4)]}
        with self.assertRaisesRegex(ValidationError, "unreadable"):
            update_workbook(path, SMDR_HEADERS, incoming, lambda file: file.write_bytes(b"partial"))
        self.assertEqual(digest(path), before)

    def test_intervening_writer_is_detected(self):
        path = self.folder / "concurrent.xlsx"
        write_workbook(path, SMDR_HEADERS, seed_sheet("avaya"))
        incoming = {"2026-08-04": read_avaya(FIXTURES["avaya"])[date(2026, 8, 4)]}
        def other_writer(_):
            path.write_bytes(b"another writer's version")
        with self.assertRaisesRegex(ValidationError, "changed during"):
            update_workbook(path, SMDR_HEADERS, incoming, other_writer)
        self.assertEqual(path.read_bytes(), b"another writer's version")

    def test_formula_like_source_is_written_as_literal_text(self):
        path = self.csv_copy(lambda rows: rows[1].__setitem__(3, "=1+1"))
        result = run_demo("avaya", path, self.folder / "literal", NOW)
        self.assertEqual(result["status"], "PUBLISHED")
        workbook = openpyxl.load_workbook(self.folder / "literal/2026-08.xlsx")
        self.assertEqual(workbook["2026-08-04"]["D2"].value, "=1+1")
        self.assertEqual(workbook["2026-08-04"]["D2"].data_type, "s")
        workbook.close()

    def test_nas_query_cap_and_naive_timestamp_block_input(self):
        with self.assertRaisesRegex(ValidationError, "query cap"):
            read_nas(FIXTURES["nas"], result_limit=8)
        records = json.loads(FIXTURES["nas"].read_text(encoding="utf-8"))
        records[0]["UTC"] = "2026-08-04T00:00:00"
        path = self.folder / "naive.json"
        path.write_text(json.dumps(records), encoding="utf-8")
        with self.assertRaisesRegex(ValidationError, "timezone"):
            read_nas(path)

    def test_month_weekend_and_dst_boundaries(self):
        # Jan UTC+11 and July UTC+10: the same UTC clock falls on different local dates.
        self.assertEqual(local_day(datetime(2026, 1, 31, 13, 30, tzinfo=timezone.utc)), date(2026, 2, 1))
        self.assertEqual(local_day(datetime(2026, 7, 31, 13, 30, tzinfo=timezone.utc)), date(2026, 7, 31))
        before = datetime(2026, 10, 3, 15, 59, tzinfo=timezone.utc).astimezone(MELBOURNE)
        after = datetime(2026, 10, 3, 16, 0, tzinfo=timezone.utc).astimezone(MELBOURNE)
        self.assertEqual((before.hour, after.hour), (1, 3))
        groups = {date(2026, 1, 31): (("Saturday",),), date(2026, 2, 1): (("Sunday",),), date(2026, 2, 2): (("today",),)}
        completed = completed_dates(groups, datetime(2026, 2, 2, 10, tzinfo=MELBOURNE))
        self.assertEqual(set(completed), {date(2026, 1, 31), date(2026, 2, 1)})
        self.assertEqual(set(by_month(completed)), {"2026-01", "2026-02"})

    def test_cli_success_and_expected_blocked_exit(self):
        for scenario, code, status in (("success", 0, "PUBLISHED"), ("blocked", 2, "BLOCKED")):
            with self.subTest(scenario=scenario):
                result = subprocess.run([sys.executable, "-m", "demo", "--workflow", "all", "--scenario", scenario,
                    "--output-dir", str(self.folder / scenario)], cwd=ROOT, capture_output=True, text=True)
                self.assertEqual(result.returncode, code, result.stderr)
                self.assertEqual(result.stdout.count(status), 2)


if __name__ == "__main__":
    unittest.main()
